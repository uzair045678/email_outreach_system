#!/usr/bin/env python3
"""
Sync approvals from drafts_for_review.csv to outreach.xlsx
Run this after you edit drafts_for_review.csv and set approve=y
"""

import pandas as pd
import openpyxl
import sys
from datetime import datetime, timedelta
import pytz

OUTREACH_XLSX = r"F:\_PROJECTS\EMails\outreach.xlsx"
REVIEW_CSV = r"F:\_PROJECTS\EMails\drafts_for_review.csv"


def load_workbook_with_retry(path):
    for attempt in range(5):
        try:
            return openpyxl.load_workbook(path)
        except PermissionError:
            if attempt < 4:
                import time
                print(f"Excel locked, retry {attempt + 1}/5...")
                time.sleep(2)
            else:
                raise


def save_workbook_with_retry(wb, path):
    for attempt in range(5):
        try:
            wb.save(path)
            return
        except PermissionError:
            if attempt < 4:
                import time
                print(f"Excel locked on save, retry {attempt + 1}/5...")
                time.sleep(2)
            else:
                raise


def parse_utc(utc_str):
    if not utc_str:
        return None
    s = str(utc_str).strip()
    if s.endswith('Z'):
        s = s[:-1] + '+00:00'
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        try:
            return datetime.strptime(s, '%Y-%m-%dT%H:%M:%S')
        except:
            return None


def get_next_business_day_9am(tz_name):
    """If original send time has passed, return next business day 9 AM in that TZ as UTC"""
    try:
        tz = pytz.timezone(tz_name)
    except:
        return None
    
    now_local = datetime.now(tz)
    next_day = now_local + timedelta(days=1)
    while next_day.weekday() >= 5:
        next_day += timedelta(days=1)
    
    next_send = tz.localize(datetime(next_day.year, next_day.month, next_day.day, 9, 0))
    return next_send.astimezone(pytz.utc)


def main():
    print("=" * 60)
    print("Syncing approvals from drafts_for_review.csv to outreach.xlsx")
    print("=" * 60)
    
    # Load review CSV
    try:
        df = pd.read_csv(REVIEW_CSV)
    except FileNotFoundError:
        print(f"ERROR: {REVIEW_CSV} not found")
        return 1
    
    print(f"Loaded {len(df)} rows from review CSV")
    
    # Filter approved
    approved = df[df['approve'].astype(str).str.lower() == 'y']
    print(f"Approved rows: {len(approved)}")
    
    if len(approved) == 0:
        print("No approvals to sync. Exiting.")
        return 0
    
    # Load outreach.xlsx
    wb = load_workbook_with_retry(OUTREACH_XLSX)
    ws = wb['Queue']
    
    # Build id -> row mapping
    id_to_row = {}
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(row=r, column=1).value
        if qid:
            id_to_row[str(qid).strip()] = r
    
    print(f"Found {len(id_to_row)} rows in outreach.xlsx Queue")
    
    # Sync each approved row
    synced = 0
    updated = 0
    skipped = 0
    now_utc = datetime.now(pytz.UTC)
    
    for _, row in approved.iterrows():
        qid = str(row['id']).strip()
        resume_var = str(row['resume_variant']).strip() if pd.notna(row['resume_variant']) else 'general'
        
        if qid not in id_to_row:
            print(f"  SKIP: {qid} not found in outreach.xlsx")
            skipped += 1
            continue
        
        r = id_to_row[qid]
        
        # Check current status
        current_status = ws.cell(row=r, column=19).value  # status
        current_approve = ws.cell(row=r, column=20).value  # approve
        
        if current_status == 'sent':
            print(f"  SKIP: {qid} ({row['company']}) already sent")
            skipped += 1
            continue
        
        # Update approve and resume_variant
        ws.cell(row=r, column=20, value='y')      # approve
        ws.cell(row=r, column=16, value=resume_var)  # resume_variant
        
        # If send_after_utc has passed, reschedule to next business day 9 AM
        send_after = ws.cell(row=r, column=21).value  # send_after_utc
        tz = ws.cell(row=r, column=23).value  # recipient_tz
        if send_after and tz:
            send_utc = parse_utc(send_after)
            if send_utc and send_utc <= now_utc:
                new_send = get_next_business_day_9am(tz)
                if new_send:
                    ws.cell(row=r, column=21, value=new_send.isoformat().replace('+00:00', 'Z'))
                    ws.cell(row=r, column=22, value=f"{new_send.astimezone(pytz.timezone(tz)).strftime('%Y-%m-%d %H:%M')} ({tz})")
                    print(f"  RESCHEDULED: {row['company']} -> next business day 9 AM {tz}")
        
        synced += 1
        updated += 1
        print(f"  SYNCED: {row['company']} ({row['to_email']}) -> approve=y, resume_variant={resume_var}")
    
    if updated > 0:
        save_workbook_with_retry(wb, OUTREACH_XLSX)
        print(f"\nSaved outreach.xlsx with {updated} updates")
    
    print(f"\n=== SUMMARY ===")
    print(f"Total approved in CSV: {len(approved)}")
    print(f"Synced to outreach.xlsx: {synced}")
    print(f"Skipped (not found/already sent): {skipped}")
    
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as e:
        import traceback
        print(f"FATAL ERROR: {e}")
        traceback.print_exc()
        sys.exit(1)