"""Storage abstraction + Excel implementation.

``Store`` is the interface the rest of the app talks to; ``XlsxStore`` is the
openpyxl-backed implementation. A SQLite or Google-Sheets backend could be
dropped in later without touching the command code.

Guarantees:
  * a rolling backup is taken before every write (keep last 14)
  * writes retry on ``PermissionError`` (Excel holding the file open)
  * if writes keep failing the changes are journaled to
    ``pending_updates.jsonl`` and replayed on the next run
"""

from __future__ import annotations

import abc
import json
import shutil
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from . import util
from . import workbook as wb_schema

QUEUE = "Queue"
CONTACTED = "Contacted"
BAD = "BadAddresses"
SUPPRESSED = "Suppressed"
LOG = "Log"
SUMMARY = "Summary"


class Store(abc.ABC):
    """Interface for the single source of truth."""

    @abc.abstractmethod
    def load(self) -> None: ...

    @abc.abstractmethod
    def save(self) -> bool: ...

    @abc.abstractmethod
    def queue_all(self) -> list[dict]: ...

    @abc.abstractmethod
    def queue_get(self, row_id: str) -> dict | None: ...

    @abc.abstractmethod
    def queue_add(self, data: dict) -> str: ...

    @abc.abstractmethod
    def queue_update(self, row_id: str, **fields: Any) -> bool: ...

    @abc.abstractmethod
    def log(self, command: str, level: str, message: str) -> None: ...


class XlsxStore(Store):
    def __init__(
        self,
        workbook_path: str | Path,
        backup_dir: str | Path,
        pending_path: str | Path,
        keep_backups: int = 14,
        logger=None,
    ):
        self.path = Path(workbook_path)
        self.backup_dir = Path(backup_dir)
        self.pending_path = Path(pending_path)
        self.keep_backups = keep_backups
        self.logger = logger
        self.wb = None
        self.pending_ops: list[dict] = []
        self._headers: dict[str, list[str]] = {}

    # -- lifecycle ----------------------------------------------------------
    def load(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)

        if self.path.exists() and self.path.stat().st_size > 0:
            self.wb = load_workbook(self.path)
        else:
            self.wb = wb_schema.build_workbook()

        changed = wb_schema.ensure_schema(self.wb)
        self._cache_headers()

        journaled = self._read_journal()
        if journaled:
            for op in journaled:
                self._apply_op(op)
            self.pending_ops = list(journaled)
            wb_schema.restyle_all(self.wb)
            self._cache_headers()
            if self.logger:
                self.logger.info("Replayed %d journaled update(s).", len(journaled))
            self.save()

        if changed:
            wb_schema.restyle_all(self.wb)
            self._cache_headers()

    def _cache_headers(self) -> None:
        self._headers = {}
        for name in wb_schema.SHEETS:
            if name in self.wb.sheetnames:
                ws = self.wb[name]
                self._headers[name] = [
                    c.value for c in ws[1] if c.value is not None
                ]

    def _ws(self, sheet: str):
        return self.wb[sheet]

    def _col(self, sheet: str, field: str) -> int:
        headers = self._headers[sheet]
        if field not in headers:
            # add the column on the fly
            ws = self._ws(sheet)
            ws.cell(row=1, column=ws.max_column + 1, value=field)
            headers.append(field)
        return headers.index(field) + 1

    # -- save / backup / journal -------------------------------------------
    def _backup(self) -> None:
        if not self.path.exists() or self.path.stat().st_size == 0:
            return
        stamp = util.now_utc().strftime("%Y%m%d-%H%M")
        dest = self.backup_dir / f"outreach-{stamp}.xlsx"
        try:
            shutil.copy2(self.path, dest)
        except OSError:
            return
        backups = sorted(self.backup_dir.glob("outreach-*.xlsx"))
        for old in backups[: -self.keep_backups]:
            try:
                old.unlink()
            except OSError:
                pass

    def save(self) -> bool:
        self._backup()
        tmp = self.path.with_suffix(".xlsx.tmp")

        def _write() -> None:
            self.wb.save(tmp)
            tmp.replace(self.path)

        try:
            util.retry(_write, attempts=5, base_delay=1.0)
        except BaseException as exc:  # noqa: BLE001
            if self.logger:
                self.logger.error(
                    "Workbook write failed (%s); journaling %d update(s) to %s",
                    exc,
                    len(self.pending_ops),
                    self.pending_path,
                )
            self._write_journal()
            return False

        self.pending_ops.clear()
        if self.pending_path.exists():
            try:
                self.pending_path.unlink()
            except OSError:
                pass
        return True

    def _write_journal(self) -> None:
        try:
            self.pending_path.parent.mkdir(parents=True, exist_ok=True)
            with self.pending_path.open("w", encoding="utf-8") as fh:
                for op in self.pending_ops:
                    fh.write(json.dumps(op, ensure_ascii=False, default=str) + "\n")
        except OSError:
            pass

    def _read_journal(self) -> list[dict]:
        if not self.pending_path.exists():
            return []
        ops: list[dict] = []
        with self.pending_path.open("r", encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if line:
                    try:
                        ops.append(json.loads(line))
                    except json.JSONDecodeError:
                        continue
        return ops

    # -- generic row helpers ------------------------------------------------
    def _iter_rows(self, sheet: str) -> Iterable[tuple[int, dict]]:
        ws = self._ws(sheet)
        headers = self._headers[sheet]
        for r in range(2, ws.max_row + 1):
            values = [ws.cell(row=r, column=c + 1).value for c in range(len(headers))]
            if all(v is None for v in values):
                continue
            yield r, dict(zip(headers, values))

    def _find_row(self, sheet: str, field: str, value: Any) -> int | None:
        col = self._col(sheet, field)
        ws = self._ws(sheet)
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=col).value == value:
                return r
        return None

    def _append_row(self, sheet: str, data: dict) -> int:
        ws = self._ws(sheet)
        headers = self._headers[sheet]
        r = ws.max_row + 1
        for field, value in data.items():
            if field not in headers:
                self._col(sheet, field)
                headers = self._headers[sheet]
            ws.cell(row=r, column=headers.index(field) + 1, value=value)
        return r

    def _set_fields(self, sheet: str, row: int, fields: dict) -> None:
        ws = self._ws(sheet)
        for field, value in fields.items():
            col = self._col(sheet, field)
            ws.cell(row=row, column=col, value=value)

    def _record(self, op: dict) -> None:
        self.pending_ops.append(op)

    def _apply_op(self, op: dict) -> None:
        sheet = op["sheet"]
        if op["kind"] == "append":
            if op.get("dedupe_key"):
                existing = self._find_row(sheet, op["dedupe_key"], op["values"].get(op["dedupe_key"]))
                if existing is not None:
                    return
            self._append_row(sheet, op["values"])
        elif op["kind"] == "upsert":
            row = self._find_row(sheet, op["key"], op["match"])
            if row is None:
                self._append_row(sheet, {op["key"]: op["match"], **op["values"]})
            else:
                self._set_fields(sheet, row, op["values"])

    # -- Queue --------------------------------------------------------------
    def queue_all(self) -> list[dict]:
        return [row for _, row in self._iter_rows(QUEUE)]

    def queue_get(self, row_id: str) -> dict | None:
        row = self._find_row(QUEUE, "id", row_id)
        if row is None:
            return None
        ws = self._ws(QUEUE)
        headers = self._headers[QUEUE]
        return {h: ws.cell(row=row, column=i + 1).value for i, h in enumerate(headers)}

    def queue_by_status(self, *statuses: str) -> list[dict]:
        wanted = set(statuses)
        return [r for r in self.queue_all() if (r.get("status") or "") in wanted]

    def queue_add(self, data: dict) -> str:
        row_id = data.get("id") or util.new_id("q_")
        data = {**data, "id": row_id}
        self._append_row(QUEUE, data)
        self._record({"sheet": QUEUE, "kind": "upsert", "key": "id", "match": row_id,
                      "values": {**data}})
        return row_id

    def queue_update(self, row_id: str, **fields: Any) -> bool:
        row = self._find_row(QUEUE, "id", row_id)
        if row is None:
            return False
        self._set_fields(QUEUE, row, fields)
        self._record({"sheet": QUEUE, "kind": "upsert", "key": "id", "match": row_id,
                      "values": fields})
        return True

    def queue_exists_domain(self, domain: str) -> bool:
        domain = util.registrable_domain(domain)
        for r in self.queue_all():
            if util.registrable_domain(util.domain_of(r.get("website") or r.get("to_email") or "")) == domain:
                return True
        return False

    # -- Contacted ----------------------------------------------------------
    def contacted_add(self, data: dict) -> None:
        data.setdefault("first_contacted", util.to_iso(util.now_utc()))
        self._append_row(CONTACTED, data)
        self._record({"sheet": CONTACTED, "kind": "append", "dedupe_key": "email",
                      "values": data})

    def contacted_emails(self) -> set[str]:
        return {str(r.get("email") or "").lower() for _, r in self._iter_rows(CONTACTED) if r.get("email")}

    def contacted_domains(self) -> set[str]:
        return {
            str(r.get("domain") or "").lower()
            for _, r in self._iter_rows(CONTACTED)
            if r.get("domain")
        }

    # -- BadAddresses -------------------------------------------------------
    def bad_add(self, data: dict) -> None:
        data.setdefault("added_at", util.to_iso(util.now_utc()))
        self._append_row(BAD, data)
        self._record({"sheet": BAD, "kind": "append", "dedupe_key": "email", "values": data})

    def bad_emails(self) -> set[str]:
        return {str(r.get("email") or "").lower() for _, r in self._iter_rows(BAD) if r.get("email")}

    def bad_domains(self) -> set[str]:
        return {str(r.get("domain") or "").lower() for _, r in self._iter_rows(BAD) if r.get("domain")}

    # -- Suppressed ---------------------------------------------------------
    def suppressed_add(self, data: dict) -> None:
        data.setdefault("added_at", util.to_iso(util.now_utc()))
        self._append_row(SUPPRESSED, data)
        self._record({"sheet": SUPPRESSED, "kind": "append", "dedupe_key": "email",
                      "values": data})

    def suppressed_emails(self) -> set[str]:
        return {str(r.get("email") or "").lower() for _, r in self._iter_rows(SUPPRESSED) if r.get("email")}

    def suppressed_domains(self) -> set[str]:
        return {str(r.get("domain") or "").lower() for _, r in self._iter_rows(SUPPRESSED) if r.get("domain")}

    # -- Log ----------------------------------------------------------------
    def log(self, command: str, level: str, message: str) -> None:
        data = {
            "timestamp": util.to_iso(util.now_utc()),
            "command": command,
            "level": level,
            "message": message[:2000],
        }
        self._append_row(LOG, data)
        self._record({"sheet": LOG, "kind": "append", "values": data})

    # -- Summary ------------------------------------------------------------
    def summary_get(self, date: str) -> dict | None:
        row = self._find_row(SUMMARY, "date", date)
        if row is None:
            return None
        ws = self._ws(SUMMARY)
        headers = self._headers[SUMMARY]
        return {h: ws.cell(row=row, column=i + 1).value for i, h in enumerate(headers)}

    def summary_bump(self, date: str, field: str, amount: int = 1) -> None:
        row = self._find_row(SUMMARY, "date", date)
        ws = self._ws(SUMMARY)
        col = self._col(SUMMARY, field)
        if row is None:
            self._append_row(SUMMARY, {"date": date, field: amount})
            self._record({"sheet": SUMMARY, "kind": "upsert", "key": "date", "match": date,
                          "values": {field: amount}})
            return
        current = ws.cell(row=row, column=col).value or 0
        ws.cell(row=row, column=col, value=current + amount)
        self._record({"sheet": SUMMARY, "kind": "upsert", "key": "date", "match": date,
                      "values": {field: current + amount}})

    def summary_set(self, date: str, **fields: Any) -> None:
        row = self._find_row(SUMMARY, "date", date)
        if row is None:
            self._append_row(SUMMARY, {"date": date, **fields})
        else:
            self._set_fields(SUMMARY, row, fields)
        self._record({"sheet": SUMMARY, "kind": "upsert", "key": "date", "match": date,
                      "values": fields})

    def today_str(self) -> str:
        return util.now_utc().strftime("%Y-%m-%d")


def open_store(cfg, logger=None) -> XlsxStore:
    store = XlsxStore(
        workbook_path=cfg.path("workbook"),
        backup_dir=cfg.path("backup_dir"),
        pending_path=cfg.path("pending_updates"),
        logger=logger,
    )
    store.load()
    return store
